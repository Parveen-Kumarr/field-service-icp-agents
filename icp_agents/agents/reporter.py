"""Agent 5 - Reporter: turns the team's work into deliverables.

Writes the ICP workbook (accounts, contacts + openers, evidence with sources,
the full agent conversation, pages scraped with fetch proof, usage), a
readable conversation transcript, and machine-readable JSON.
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..vendor_profile import RUBRIC, TIERS, rubric_text
from ..messages import AgentMessage
from ..scraper import now_iso
from .base import BaseAgent

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3A5F")
TIER_FILL = {"A": "D9EAD3", "B": "FFF2CC", "C": "FCE5CD"}


def _tier_key(t: str) -> str:
    return t[:1] if t[:1] in "ABC" and t[1:3] == " -" else "Z"


class ReporterAgent(BaseAgent):
    name = "Reporter"
    role = "Writes the ICP workbook, the conversation transcript and JSON results"

    def __init__(self, runtime, llm=None, settings=None):
        super().__init__(runtime, llm, settings)
        self.paths: dict[str, str] = {}

    async def handle(self, msg: AgentMessage) -> None:
        if msg.subject == "write_report":
            paths = self.write()
            await self.reply(msg, f"Done. Workbook: {paths['xlsx']}. Transcript: {paths['transcript']}. "
                                  f"JSON: {paths['json']}.", paths)

    # --- gather ---------------------------------------------------------------------
    def rows(self) -> tuple[list[dict], list[dict], list[dict]]:
        ag = self.runtime.agents
        scout, researcher, analyst, strategist = ag["Scout"], ag["Researcher"], ag["Analyst"], ag["Strategist"]
        coord = ag["Coordinator"]
        accounts, contacts, evidence = [], [], []
        for acct_id in coord.expected or []:
            a = scout.accounts.get(acct_id, {"company": acct_id, "roles": [], "contacts": []})
            r = researcher.records.get(acct_id, {})
            v = analyst.verdicts.get(acct_id, {})
            p = strategist.plays.get(acct_id, {})
            st = coord.status.get(acct_id, {})
            accounts.append({
                "account": r.get("canonical_name") or a["company"],
                "tier": v.get("tier") or ("Error: " + str(st.get("error", "incomplete"))[:80]),
                "total": v.get("total"), "scores": v.get("scores", {}),
                "industry": r.get("industry", ""), "what_they_service": r.get("what_they_service", ""),
                "scale": r.get("scale", ""), "size_evidence": r.get("size_evidence", ""),
                "service_org": r.get("service_organisation", ""),
                "stack": ", ".join(r.get("fsm_stack_signals") or []),
                "ai": ", ".join(r.get("ai_or_digital_initiatives") or []),
                "relationship": r.get("relationship_to_vendor", ""), "roles": ", ".join(a.get("roles", [])),
                "speakers": len(a.get("contacts", [])), "summary": v.get("summary", ""),
                "questions": len(v.get("questions_asked", [])) + len(p.get("dialogue", [])),
                "lead_with": ", ".join(p.get("lead_with", [])), "why_now": p.get("why_now", ""),
                "next_step": p.get("next_step", ""), "confidence": r.get("confidence", ""),
                "sources": len(r.get("sources", [])), "website": r.get("website", ""),
            })
            openers = {o.get("contact", "").lower(): o for o in p.get("openers", [])}
            for c in a.get("contacts", []) or [{"name": "", "title": ""}]:
                o = openers.get(c["name"].lower()) or (next(iter(openers.values())) if not c["name"] and openers else {})
                if c["name"] or o:
                    contacts.append({"name": c["name"] or "(role to find)", "title": c.get("title", ""),
                                     "account": accounts[-1]["account"], "tier": accounts[-1]["tier"],
                                     "total": v.get("total"), "channel": o.get("channel", ""),
                                     "opener": o.get("message", ""), "next_step": p.get("next_step", "")})
            for f in r.get("facts") or []:
                if isinstance(f, str):
                    f = {"claim": f}
                if isinstance(f, dict):
                    evidence.append({"account": accounts[-1]["account"], "claim": str(f.get("claim", "")),
                                     "source": str(f.get("source_url", ""))})
        accounts.sort(key=lambda x: (_tier_key(x["tier"]), -(x["total"] or 0), x["account"]))
        contacts.sort(key=lambda x: (_tier_key(x["tier"]), -(x["total"] or 0), x["account"]))
        return accounts, contacts, evidence

    # --- write ----------------------------------------------------------------------
    def write_partial(self) -> str:
        """Overwrite a 'so far' workbook + transcript; called after every finished account."""
        out = Path(self.settings.output_dir)
        out.mkdir(parents=True, exist_ok=True)
        accounts, contacts, evidence = self.rows()
        path = out / "FSN_West_ICP_partial.xlsx"
        self._xlsx(str(path), accounts, contacts, evidence)
        self._transcript(str(out / "agent_conversation_partial.md"))
        return str(path)

    def write(self) -> dict[str, str]:
        out = Path(self.settings.output_dir)
        out.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        accounts, contacts, evidence = self.rows()
        paths = {"xlsx": str(out / f"FSN_West_ICP_{stamp}.xlsx"), "transcript": str(out / f"agent_conversation_{stamp}.md"),
                 "json": str(out / f"icp_results_{stamp}.json")}
        self._xlsx(paths["xlsx"], accounts, contacts, evidence)
        self._transcript(paths["transcript"])
        coord = self.runtime.agents["Coordinator"]
        Path(paths["json"]).write_text(json.dumps({
            "run": self._run_info(), "debrief": coord.debrief, "accounts": accounts, "contacts": contacts,
            "evidence": evidence, "conversation": [m.to_record() for m in self.runtime.transcript],
        }, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
        self.paths = paths
        return paths

    def _run_info(self) -> dict:
        coord = self.runtime.agents["Coordinator"]
        scout = self.runtime.agents["Scout"]
        llm = self.llm
        usage = {k: vars(u) for k, u in (llm.usage.items() if llm else [])}
        return {"event_url": self.settings.event_url, "model": getattr(llm, "model", None),
                "finished_at": now_iso(), "pages": [p.record() for p in scout.pages],
                "capture": str(scout.capture_path), "usage": usage, "conversation": self.runtime.stats(),
                "accounts_found": coord.scrape_info.get("total_found"),
                "speakers_found": coord.scrape_info.get("speakers")}

    def _transcript(self, path: str) -> None:
        lines = ["# Agent conversation", ""]
        for m in self.runtime.transcript:
            reply = f" (reply to #{m.in_reply_to})" if m.in_reply_to else ""
            lines += [f"**#{m.id} {m.ts.strftime('%H:%M:%S')} {m.sender} -> {m.recipient}** "
                      f"`{m.performative}` _{m.subject}_{reply}", "", m.text, ""]
        Path(path).write_text("\n".join(lines), encoding="utf-8")

    def _sheet(self, wb, title, headers, rows, widths, name, wrap=()):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.font = Font(name=FONT, bold=True, color="FFFFFF")
            c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name=FONT, size=10)
                c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap)
        ws.freeze_panes = "B2"
        tab = Table(displayName=name, ref=f"A1:{get_column_letter(len(headers))}{max(2, len(rows) + 1)}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tab)
        return ws

    def _xlsx(self, path, accounts, contacts, evidence) -> None:
        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        dims = list(RUBRIC)

        acc_rows = []
        for n, a in enumerate(accounts, 2):
            pts = [a["scores"].get(d, {}).get("points") for d in dims]
            first, last = get_column_letter(4), get_column_letter(3 + len(dims))
            total = f"=SUM({first}{n}:{last}{n})" if a["total"] is not None else None
            acc_rows.append([a["account"], a["tier"], total, *pts, a["summary"], a["industry"], a["what_they_service"],
                             a["scale"], a["size_evidence"], a["service_org"], a["stack"], a["ai"], a["relationship"],
                             a["roles"], a["speakers"], a["lead_with"], a["why_now"], a["next_step"], a["questions"],
                             a["confidence"], a["sources"], a["website"]])
        ws = self._sheet(wb, "Accounts (ICP)",
                         ["Account", "ICP tier", "Score (0-100)", *[f"{d.title()} (/{RUBRIC[d][0]})" for d in dims],
                          "Analyst's reasoning", "Industry", "What they service", "Scale", "Size evidence",
                          "Service organisation", "FSM/CRM stack signals", "AI / digital initiatives",
                          "Relationship to vendor", "Seen at event as", "# Speakers", "Lead with (vendor capabilities)",
                          "Why now", "Next step", "Agent questions on this account", "Research confidence",
                          "# Sources", "Website"],
                         acc_rows, [26, 28, 10, 10, 10, 10, 10, 10, 60, 22, 40, 11, 36, 40, 30, 30, 18, 24, 10, 34,
                                    44, 40, 12, 12, 9, 28], "Accounts", wrap=(9, 11, 13, 14, 15, 16, 22, 23, 24))
        for i, a in enumerate(accounts, 2):
            fill = TIER_FILL.get(_tier_key(a["tier"]))
            if fill:
                ws.cell(i, 2).fill = PatternFill("solid", fgColor=fill)

        self._sheet(wb, "Contacts & Openers", ["Name", "Title", "Account", "ICP tier", "Score", "Channel",
                                               "Opener draft (human reviews and sends)", "Next step"],
                    [[c["name"], c["title"], c["account"], c["tier"], c["total"], c["channel"], c["opener"],
                      c["next_step"]] for c in contacts],
                    [22, 34, 26, 26, 8, 11, 90, 44], "Contacts", wrap=(2, 7, 8))
        self._sheet(wb, "Evidence", ["Account", "Claim", "Source"],
                    [[e["account"], e["claim"], e["source"]] for e in evidence], [26, 90, 60], "Evidence", wrap=(2,))
        self._sheet(wb, "Agent Conversation", ["#", "Time (UTC)", "From", "To", "Type", "Subject", "In reply to", "Message"],
                    [[m.id, m.ts.strftime("%H:%M:%S"), m.sender, m.recipient, m.performative, m.subject,
                      m.in_reply_to, m.text] for m in self.runtime.transcript],
                    [6, 11, 12, 12, 10, 20, 10, 120], "Conversation", wrap=(8,))
        scout = self.runtime.agents["Scout"]
        self._sheet(wb, "Pages Scraped (live)", ["Page", "URL", "HTTP", "Channel", "Fetched at (UTC)", "Bytes",
                                                 "SHA-256 (16)", "Rendered JS", "Error"],
                    [[p.role, p.url, p.status, p.channel, p.fetched_at, p.bytes, p.sha256, p.rendered_js, p.error]
                     for p in scout.pages], [14, 60, 7, 16, 24, 10, 20, 11, 40], "Pages")
        if self.llm:
            self._sheet(wb, "Usage", ["Agent", "Model calls", "Input tokens", "Output tokens", "Thinking tokens",
                                      "Web searches", "Pages read"],
                        [[k, u.calls, u.input_tokens, u.output_tokens, u.thought_tokens, u.web_searches, u.url_fetches]
                         for k, u in self.llm.usage.items()], [14, 12, 14, 14, 14, 15, 11], "Usage")
        if hasattr(self.llm, "backends"):
            self._sheet(wb, "Model Pool", ["Provider", "Model", "Calls", "Input tokens", "Output tokens",
                                           "Rate-limited (429)", "Errors", "Avg seconds/call", "Status"],
                        [[b.name, b.model, b.stats["calls"], b.stats["in_tokens"], b.stats["out_tokens"],
                          b.stats["rate_limited"], b.stats["errors"],
                          round(b.stats["seconds"] / b.stats["calls"], 1) if b.stats["calls"] else None,
                          b.disabled or b.exhausted or "ok"] for b in self.llm.backends],
                        [12, 30, 8, 13, 13, 12, 8, 12, 50], "Pool")
        self._sheet(wb, "ICP Rubric", ["Dimension", "Max", "How it is scored"],
                    [[k, mx, d] for k, (mx, d) in RUBRIC.items()]
                    + [["Tiers", "", rubric_text().splitlines()[-1]]], [16, 8, 120], "Rubric", wrap=(3,))

        # Summary
        s = summary
        s.column_dimensions["A"].width = 34
        s.column_dimensions["B"].width = 100
        coord = self.runtime.agents["Coordinator"]
        info = self._run_info()
        stats = info["conversation"]
        s["A1"] = "Field Service Next West - live ICP validation by the agent team"
        s["A1"].font = Font(name=FONT, bold=True, size=14)
        rows = [("Event page", info["event_url"]), ("Vendor (ICP owner)", self.settings.vendor_name or "(generic field-service AI vendor profile)"), ("Run finished (UTC)", info["finished_at"]),
                ("Data source", "Live scrape at run time (see 'Pages Scraped (live)' for HTTP status, time and hash)"),
                ("Companies found live", info["accounts_found"]), ("Speakers found live", info["speakers_found"]),
                ("Accounts processed", len(accounts)), ("Model", info["model"]),
                ("Agent messages", stats["messages"]), ("Questions asked and answered between agents",
                                                         stats["questions_answered"])]
        r = 3
        for k, v in rows:
            s.cell(r, 1, k).font = Font(name=FONT, bold=True)
            s.cell(r, 2, v).font = Font(name=FONT)
            s.cell(r, 2).alignment = Alignment(horizontal="left")
            r += 1
        r += 1
        s.cell(r, 1, "Accounts by tier").font = Font(name=FONT, bold=True, size=12)
        r += 1
        rng = f"'Accounts (ICP)'!$B$2:$B${len(accounts) + 1}"
        for label, crit in [(lbl, lbl[:1] + " - *") for _, lbl in TIERS] + [
                ("Not ICP", "Not ICP*"), ("Competitors", "Competitor*"), ("Vendors / partners", "Vendor*"),
                ("Consultants", "Influencer*"), ("Self (the vendor itself)", "Self*"), ("Needs review", "Needs review*"),
                ("Errors", "Error*")]:
            s.cell(r, 1, label).font = Font(name=FONT)
            s.cell(r, 2, f'=COUNTIF({rng},"{crit}")').font = Font(name=FONT)
            s.cell(r, 2).alignment = Alignment(horizontal="left")
            r += 1
        s.cell(r, 1, "Total accounts").font = Font(name=FONT, bold=True)
        s.cell(r, 2, f"=COUNTA({rng})").alignment = Alignment(horizontal="left")
        r += 2
        d = coord.debrief or {}
        if d:
            s.cell(r, 1, "Coordinator's debrief").font = Font(name=FONT, bold=True, size=12)
            r += 1
            for label, val in [("Headline", d.get("headline", ""))] + \
                    [("Top account", f"{t['account']}: {t['why']}") for t in d.get("top_accounts", [])] + \
                    [("Pattern", p) for p in d.get("patterns", [])] + \
                    [("Recommended action", x) for x in d.get("recommended_actions", [])]:
                s.cell(r, 1, label).font = Font(name=FONT)
                c = s.cell(r, 2, val)
                c.font = Font(name=FONT)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                r += 1
        wb.save(path)
