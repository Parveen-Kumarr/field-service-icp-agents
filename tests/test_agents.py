"""Tests for the live agent team. Run: python -m pytest -q

No API key and no network are needed: Gemini is replaced by tests/fake_gemini.py
and the conference site by an in-memory HTTP transport serving pages that follow
the real site's structure. The runtime, tool loop, scraper and agents are the
real code.
"""
import asyncio
import sys
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from fake_gemini import FakeGemini  # noqa: E402
from icp_agents import LLM, Settings, run_team  # noqa: E402
from icp_agents import scraper  # noqa: E402
from icp_agents.messages import REPLY, REQUEST  # noqa: E402

FIX = ROOT / "tests" / "fixtures"
BASE = "https://fieldserviceusa.wbresearch.com/"


def site(down: bool = False):
    pages = {"/": "home.html", "/speakers": "speakers.html", "/sponsors": "sponsors.html"}

    def handler(request: httpx.Request) -> httpx.Response:
        if down:
            return httpx.Response(503)
        if request.url.path == "/attendee-list":
            return httpx.Response(405)
        if request.url.path.startswith("/media/logos/"):  # logo images, read by Gemini vision
            return httpx.Response(200, content=f"PNG:{request.url.path}".encode(), headers={"content-type": "image/png"})
        f = pages.get(request.url.path)
        return httpx.Response(200, text=(FIX / f).read_text()) if f else httpx.Response(404)

    return scraper.LiveFetcher("test-agent", client=httpx.AsyncClient(transport=httpx.MockTransport(handler)))


def settings(tmp_path, **kw) -> Settings:
    s = Settings()
    s.event_url, s.output_dir, s.render_js, s.concurrency, s.max_accounts = BASE, str(tmp_path), False, 3, 0
    s.vendor_name, s.vendor_url, s.heartbeat = "Acme Service AI", "https://acme-service-ai.example/", 0.05
    s.fast = False  # full research mode unless a test asks for fast mode
    for k, v in kw.items():
        setattr(s, k, v)
    return s


def run(tmp_path, down=False, **kw):
    llm = LLM(model="gemini-test", client=FakeGemini())
    rt, team = asyncio.run(run_team(settings(tmp_path, **kw), llm, fetcher=site(down), timeout=60))
    return rt, team, llm


# --- scraper -------------------------------------------------------------------------
def test_speaker_cards_parse():
    got = scraper.parse_speakers((FIX / "speakers.html").read_text())
    assert [g["name"] for g in got][:2] == ["Clint Olsen", "Greg Friesen"]
    assert got[1]["title"] == "VP & General Manager, Global Services" and got[1]["company"] == "Ciena"
    assert len(got) == 6  # company logos ("Stryker Logo") are not people


def test_pages_discovered_from_live_navigation():
    found = scraper.discover_pages((FIX / "home.html").read_text(), BASE)
    assert found["speakers"] == BASE + "speakers"
    assert found["sponsors"] == BASE + "sponsors"  # not "Sponsorship Opportunities"
    assert found["attendees"] == BASE + "attendee-list"
    assert "2024" not in " ".join(found.values())


def test_logo_strip_found_and_social_icons_ignored():
    groups = scraper.pick_logo_sections(scraper.logo_groups((FIX / "home.html").read_text(), BASE), min_logos=2)
    [(heading, logos)] = groups.items()
    assert heading.startswith("Leading the Way")
    assert len(logos) == 3 and not any("linkedin" in l["src"] for l in logos)


# --- the team, end to end ---------------------------------------------------------------
def test_live_team_run_end_to_end(tmp_path):
    rt, team, llm = run(tmp_path)
    coord, scout = team["coordinator"], team["scout"]
    assert coord.failed is None and coord.report and Path(coord.report["xlsx"]).exists()

    # Live collection: every page fetched now, with proof; gated attendee list reported, not faked
    pages = {p.role: p for p in scout.pages}
    assert pages["home"].status == 200 and pages["home"].sha256 and pages["home"].fetched_at
    assert pages["attendees"].status == 405
    assert any("gated" in m.text for m in rt.transcript if m.sender == "Scout")
    assert set(coord.expected) >= {"stryker", "ciena", "crown equipment", "aquant", "sam service", "tetra pak",
                                    "omron", "dell technologies", "acme service ai", "ifs"}
    crown = scout.accounts["crown equipment"]
    assert len(crown["contacts"]) == 2 and len(crown["name_variants"]) == 2  # "Corp." variant merged

    # The Researcher read the configured vendor site and briefed the team before research
    brief = next(m for m in rt.transcript if m.subject == "vendor_briefing")
    assert brief.recipient == "*" and team["analyst"].briefing["source"] == "live"

    # Real questions and answers: the Analyst asked about the low-confidence account and waited
    q = next(m for m in rt.transcript if m.performative == REQUEST and m.sender == "Analyst" and m.thread == "sam service")
    a = next(m for m in rt.transcript if m.performative == REPLY and m.in_reply_to == q.id)
    assert a.sender == "Researcher" and "mechanical services contractor" in a.text
    assert q.ts <= a.ts
    assert team["analyst"].verdicts["sam service"]["questions_asked"][0]["a"] == a.text

    # The Strategist asked the Researcher for a hook on every play, and challenged one verdict
    plays = team["strategist"].plays
    assert plays and all(any(d["with"] == "Researcher" for d in p["dialogue"]) for p in plays.values())
    ch = next(m for m in rt.transcript if m.sender == "Strategist" and m.data.get("challenge"))
    ans = next(m for m in rt.transcript if m.in_reply_to == ch.id)
    assert "revised" in ans.text and team["analyst"].verdicts["tetra pak"]["tier"].startswith("B")
    assert any(m.subject == "verdict_revised" for m in rt.transcript)

    # Verdicts: tier comes from code-summed scores; disqualifiers respected
    v = team["analyst"].verdicts
    assert v["stryker"]["tier"].startswith("A") and v["stryker"]["total"] == 92
    assert v["aquant"]["tier"].startswith("Competitor") and v["acme service ai"]["tier"].startswith("Self")
    assert v["dell technologies"]["tier"] == "B - Good ICP (also a vendor/sponsor)"
    assert v["ifs"]["tier"].startswith("Vendor")
    # only real A/B accounts get a play (never the vendor itself)
    assert set(plays) == {"stryker", "ciena", "crown equipment", "tetra pak", "omron", "dell technologies"}

    # Tool loop: prose answers were nudged into a forced submit; searches, URL reads and citations counted
    u = llm.total_usage()
    assert u.web_searches > 0 and u.url_fetches > 0 and u.thought_tokens > 0
    assert team["researcher"].records["stryker"]["sources"][0]["url"] == "https://example.com/stryker"
    # Follow-ups are stateful: they carry previous_interaction_id and only the new function results
    reqs = llm.client.interactions.requests
    follow = [r for r in reqs if r.get("previous_interaction_id")]
    assert follow and all(r["input"][0]["type"] in ("function_result", "text") for r in follow)
    assert any(r["input"][0]["type"] == "function_result" and r["input"][0]["call_id"].startswith("fc_") for r in follow)
    # Logos with no alt text were read from the downloaded images
    assert {"tetra pak", "omron"} <= set(coord.expected)

    # Workbook: conversation, pages with proof, formulas
    wb = load_workbook(coord.report["xlsx"])
    assert {"Accounts (ICP)", "Contacts & Openers", "Evidence", "Agent Conversation", "Pages Scraped (live)",
            "Usage"} <= set(wb.sheetnames)
    assert wb["Agent Conversation"].max_row - 1 == len(rt.transcript) - 1  # all but the Reporter's own reply
    assert str(wb["Accounts (ICP)"]["C2"].value).startswith("=SUM(")
    assert Path(coord.report["transcript"]).read_text().count("**#") >= 40


def test_force_live_stops_when_site_unreachable(tmp_path):
    rt, team, _ = run(tmp_path, down=True, scrape_mode="local")
    coord = team["coordinator"]
    assert coord.failed and "No live data" in coord.failed and coord.report is None
    assert any(m.subject == "abort" for m in rt.transcript)


def test_gemini_reads_site_live_when_local_http_is_blocked(tmp_path):
    rt, team, llm = run(tmp_path, down=True, scrape_mode="live")
    coord, scout = team["coordinator"], team["scout"]
    assert coord.failed is None
    assert set(coord.expected) == {"stryker", "ciena", "aquant"}
    assert any(p.channel == "gemini-url_context" for p in scout.pages)
    assert any("Google's servers" in m.text for m in rt.transcript if m.sender == "Scout")
